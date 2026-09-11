"""Release gate and audit. Fails closed when any required evidence is absent."""
import hashlib, json, re
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
from collections import Counter
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent;QA=ROOT/'qa';OUT=ROOT.parent.parent/'outputs/flicker_teo_final/output'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def js(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def number(x,n=2):return f'{x:,.{n}f}'.replace(',',' ').replace('.',',').replace('-','−')
def money(x):return number(x/1e6)
results=js(QA/'verified_results.json');tests=js(QA/'input_tests.json');visual=js(QA/'visual_review.json')
assert results['failed']==0, 'Independent validation has failing checks'
assert all(t['passed'] for t in tests), 'Input regression tests have failing checks'
assert visual['all_sheets_passed'] and visual['all_document_pages_passed'], 'Visual gate is not closed'
required=['Flicker_TEO_FINAL.xlsx','Flicker_TEO_FINAL.docx','Flicker_TEO_FINAL.pdf']
assert all((OUT/name).stat().st_size>1024 for name in required)
sources=[(Path('C:/Users/novik/Downloads/00_Flicker_Master_Context (1).md'),'53f0fc16cb5e1751d561d61cba2e3c53ae767102fe9cc1c9da211c505bb8337b'),(Path('C:/Users/novik/Downloads/Flicker_TEO_financial_model (4).xlsx'),'2e2ece6cd15b3a930db6a1384305a6ef21716c9678d7e3fc1014cfdaF65eef70'.lower())]
for p,h in sources:
    assert sha(p)==h, f'Original changed: {p}'
    assert sha(ROOT/'source'/p.name)==h, 'Source copy differs'
w=openpyxl.load_workbook(OUT/required[0],data_only=True)
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
assert not errors, errors[:10]
assert len(w.sheetnames)==22
assert set(visual['sheets'])==set(w.sheetnames)
pdf=PdfReader(OUT/required[2]);assert len(pdf.pages)==len(visual['document_pages'])
assert sorted(visual['document_pages'])==list(range(1,len(pdf.pages)+1))
assert visual['docx_sha256']==sha(OUT/required[1]) and visual['pdf_sha256']==sha(OUT/required[2]), 'Document changed after visual review'
doc=Document(OUT/required[1]);doc_text=' '.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
pdf_text=' '.join(p.extract_text() or '' for p in pdf.pages)
norm=lambda x:re.sub(r'\s+',' ',x).replace('\u00a0',' ').replace('\u202f',' ')
doc_text=norm(doc_text);pdf_text=norm(pdf_text)
numeric_pattern=re.compile(r'(?<![\w])(?:[−–-]?\d{1,3}(?: \d{3})+|[−–-]?\d+)(?:[,.]\d+)?')
def numeric_tokens(t):
    return Counter(x.replace(' ','').replace('−','-').replace('–','-').replace(',','.') for x in numeric_pattern.findall(t))
missing_numeric=numeric_tokens(doc_text)-numeric_tokens(pdf_text)
assert not missing_numeric, f'DOCX numeric tokens missing from PDF: {missing_numeric}'
base=results['base'];sc=results['scenarios'][1];mc=results['mc']
values=[('NPV',money(sc['npv'])),('Cash NPV',money(sc['cashnpv'])),('Annual benefit',money(sc['annualbenefit'])),('Annual cost',money(sc['annualcost'])),('Annual net',money(sc['annualbenefit']-sc['annualcost'])),('Prevented',number(base['count'])),('Exit cost',number(base['exitcost'])),('Pre scale',money(base['pre_scale'])),('S0 resource',money(base['phasecost'][0])),('MC P10',money(mc['p10'])),('MC P50',money(mc['p50'])),('MC P90',money(mc['p90'])),('Power',number(base['power']*100)),('Required sample',number(base['required_n'],0))]
for label,value in values:
    assert value in doc_text, f'{label} absent from DOCX: {value}'
    assert value in pdf_text, f'{label} absent from PDF: {value}'
for row,key in [(7,'npv'),(8,'cashnpv'),(13,'annualbenefit'),(14,'annualcost')]:assert abs(w['Сценарии'].cell(row,4).value-sc[key])<.02
assert all('неизвестно' in str(w['Вводные'].cell(57,c).value) for c in range(4,8))
assert w['Сценарии']['C30'].value=='не рассчитано'
assert len(doc.inline_shapes)>=10

old=js(QA/'model_data.json')['old_audit']
rows=['| Контрольный показатель | Исходный кэш | Независимый расчёт | Разница |','|---|---:|---:|---:|']
for a in old:rows.append(f'| {a["metric"]} | {number(a["cached"],6)} | {number(a["independent"],6)} | {number(a["difference"],9)} |')
bridges=[('Годовая эксплуатация',34255682.979985654,sc['annualcost'],'Взносы с бонусом и травматизм; отдельное время участников и руководителей; резерв'),('Стоимость одного ухода',1248004.761904762,base['exitcost'],'Вычет избежимой оплаты вакансии и отдельные часы рекрутера; открытые знания и юридические расходы не обнулены'),('Годовой эффект',14793948.288,sc['annualbenefit'],'Исправлена монетизация ухода; базовая воронка сохранена'),('До масштабирования',42813513.20537661,base['pre_scale'],'Добавлена подготовка, 12 месяцев валидации, внешняя экспертиза и время участников'),('Годовой чистый эффект',-19461734.691985656,sc['annualbenefit']-sc['annualcost'],'Изменились и эффект, и полная стоимость')]
bridge=['| Показатель | Было ₽ | Стало ₽ | Причина |','|---|---:|---:|---|']+[f'| {label} | {number(a)} | {number(b)} | {why} |' for label,a,b,why in bridges]
coverage=[
('Сохранность двух источников','SHA-256 оригиналов и копий, указаны ниже'),
('Полезные исходные листы и недостающая архитектура','22 листа; исходные 12 назначений сохранены и переработаны'),
('Статусы вводов и владельцы','Вводные, Запрос данных, Источники, статьи затрат; неизвестные не заменены нулями'),
('Cash, существующий ресурс и альтернативная стоимость','Команда, Прямые расходы, Время участников, Стоимость Flicker, помесячные отдельные колонки'),
('Полнота людей и расходов','13 ролевых строк, инфраструктура, экспертиза, пентест, обучение, опросы, сигналы, резерв; открытые лицензии и legal обозначены'),
('Стоимость ухода по ролям','8 групп; найм, рекрутер, интервью, вакансия и экономия, onboarding, наставник, ramp-up; неизвестные знания и legal'),
('Разделение масштаба, сигнала и причинности','Удержание; условные знаменатели; альтернативный ITT без повторной воронки'),
('Причинная научная валидация','Валидация и два раздела DOCX; outcome, cluster RCT, stepped-wedge, DiD, backtest, пропуски, exclusions, остановка, монетизация'),
('Мощность и MDE','Формулы Excel сверены со statistics.NormalDist и независимым размером выборки'),
('Четыре связных сценария','Сценарии; стоимость усиленной программы растёт; вероятности неизвестны и взвешенный NPV не рассчитан'),
('Риски и отсутствие двойного учёта','15 рисков, owners, механизмы, ячейки, валовое/остаточное влияние, expected-cost формулы, открытые хвосты'),
('Пять альтернатив','Альтернативы и DOCX; неизвестные цены, сроки, эффекты и NPV, структура запроса и пороговая оценка'),
('61 месячная точка каждого сценария','244 строки; независимый пересчёт каждого денежного и экономического потока и PV'),
('NPV, IRR, ROI, BCR и payback','Сценарии и денежный поток; IRR без смены знака не определена; отдельный тест положительной экономики'),
('One way и два фактора','7 драйверов, независимая сверка; tornado в DOCX, двухфакторная таблица и пороги в XLSX'),
('Monte Carlo','10 000 итераций, seed 20260908, редактируемые распределения, P10/50/90, вероятность, 3/5 лет, expected loss и DaR'),
('Воспроизводимость Monte Carlo','JS Park–Miller uniforms и независимый Python generator совпадают; Excel геометрические суммы сверены с месячными расчётами'),
('Этапы, stop/go и ограниченный запрос','6 этапов, сроки, команды, результаты и критерии; S0 лимит ресурса 1,80 млн ₽, cash требует подтверждения'),
('Expected cost to failure и ценность информации','Этапы; формулы и определения; численные значения не рассчитаны без неизвестных условных вероятностей и prior'),
('Полноценный DOCX и PDF','Связное обоснование, 23 тематических требований и 10 графиков; PDF экспортирован из этого DOCX'),
('Визуальный стандарт','Каждый лист и каждая страница открыты и проверены; подробное покрытие ниже'),
('Автоматический и независимый QA','Лист QA, входные regression tests, Excel пересчёт, полный скан ошибок, независимые сценарии и MC'),
('Межформатная согласованность','Ключевые значения найдены в DOCX и PDF и совпадают с финальным XLSX'),
('Независимая рекомендация','Пилот не готов до gates; масштабирование и базовая экономика не одобрены; данные, способные изменить решение, перечислены')]
coverage_md=['| Требование | Проверяемое покрытие |','|---|---|']+[f'| {a} | {b} |' for a,b in coverage]
hashes={name:sha(OUT/name) for name in required}
lines=[
'# Flicker ТЭО — аудит и validation report',
'','Версия 1.0. Дата анализа данных 08.09.2026. Финальная проверка '+datetime.now().strftime('%d.%m.%Y')+'.',
'','## Итог проверки','',
f'Пакет содержит 22 листа финансовой модели, DOCX, согласованный PDF на {len(pdf.pages)} страницах и этот аудит. Выполнено {len(results["checks"])} независимых и структурных проверок и {len(tests)} проверок поведения при изменении вводов. Ошибок формул в финальной книге не обнаружено. Пройденные проверки относятся к поставленной версии, а не к будущим редактированиям или научной эффективности Flicker.',
'','## Воспроизведение исходной модели','',*rows,
'','Восемь контрольных показателей воспроизведены до изменений. Gross независимо найден обратным подбором к прямой прогрессивной налоговой функции, а не копированием исходной обратной формулы. Допуск финансовой сверки 0,01 ₽; небольшая плавающая разница не округлялась для сокрытия расхождений.',
'','## Мост исправлений','',*bridge,
'','Исходная вероятность базового сценария 100% заменена неизвестной; нерассчитанный взвешенный результат не равен нулю. Окупаемость от начала проекта при недостижении выводится текстом, не нулём месяцев.',
'','## Базовый вывод','',
f'Годовой эффект {money(sc["annualbenefit"])} млн ₽; полная эксплуатация {money(sc["annualcost"])} млн ₽; чистый эффект {money(sc["annualbenefit"]-sc["annualcost"])} млн ₽. Экономический NPV {money(sc["npv"])} млн ₽, денежный NPV {money(sc["cashnpv"])} млн ₽. Эти значения являются условными результатами допущений. Утверждённого внутреннего бюджета и результата пилота нет.',
'','Предлагается рассмотреть только S0 с лимитом совокупного ресурса 1,80 млн ₽. Это не доказанная положительная EVSI и не автоматическое финансирование разработки. Масштабирование сейчас не рекомендовано; пилот требует privacy, power и action gates.',
'','## Воспроизводимость и численные допуски','',
'Основная книга создана средствами artifact-tool и пересчитана Microsoft Excel 16.0. Независимый Python расчёт использует прямую помесячную арифметику, в отличие от геометрических рядов sensitivity/Monte Carlo в книге. Для потоков и NPV допуск 0,02 ₽, для выборочных MC строк и квантилей 0,10 ₽. Для мощности использована независимая стандартная нормальная функция. Формулы длиннее ограничения Excel 8192 символа отсутствуют, внешних ссылок на книги нет.',
'','Все 30 000 результирующих ячеек Monte Carlo (NPV, поток 36 и 60 месяцев), все 70 000 выборок семи параметров и 25 узлов двухфакторной таблицы сверены с независимым расчётом. Числовые токены текста и таблиц DOCX дополнительно проверены на присутствие в PDF с сохранением кратности; графики проверены по исходному сверенному набору результатов.',
'','Для DOCX выполнена попытка стандартного render_docx.py; LibreOffice/soffice отсутствует в среде. Применён установленный Microsoft Word 16.0: тот же DOCX экспортирован в PDF, все страницы растрированы Poppler и просмотрены. Это проверка фактического Word-рендера, а не заявление о совместимости с LibreOffice. Для Excel устранена несовместимость рендера с некавыченными кириллическими ссылками; формулы сохранены живыми.',
f'','Monte Carlo: 10 000 испытаний, Park–Miller xₙ = 16807xₙ₋₁ mod 2147483647, seed 20260908. AF:AJ содержат фиксированные базовые целочисленные выборки; параметры распределений редактируемые и пересчитывают результаты. P10/P50/P90 NPV: '+ '/'.join(money(mc[k]) for k in ['p10','p50','p90'])+' млн ₽. Во всех выполненных испытаниях NPV отрицателен; это не нулевая вероятность успеха в действительности.',
'','## Покрытие требований','',*coverage_md,
'','## Визуальная и межформатная проверка','',
f'Все {len(w.sheetnames)} листа отрендерены и просмотрены. Для длинных таблиц и расчётных блоков использованы отдельные области; MC дополнительно проверен численно на всех 10 000 итерациях. Все {len(pdf.pages)} страницы PDF, экспортированного из поставленного DOCX, просмотрены. Исходные десять диаграмм проверены отдельно. Рендер не заменял проверку чисел.',
'','Список проверенных листов: '+', '.join(w.sheetnames)+'.',
 '','Пять редактируемых диаграмм дополнительно экспортированы из Microsoft Excel и просмотрены. У tornado подписи категорий проверены слева от полос: альтернативный preview-движок некорректно отображал это положение оси. Финальный экспорт Excel подтверждает корректную компоновку.',
'','DOCX и PDF содержат одинаковые ключевые показатели: '+', '.join(label for label,value in values)+'. Визуальная проверка включала обрезку, переносы, читаемость, легенды, единицы, отрицательные значения и отсутствие случайных пустых страниц.',
'','## Что можно утверждать перед Avito','',
'- Исходная арифметика воспроизведена, неполноты и их влияние явно показаны.\n- При указанных допущениях базовая эксплуатация и проект в целом убыточны.\n- Полная стоимость отличается от нового денежного запроса.\n- Подготовлен проверяемый план проверки данных, приватности, причинного эффекта и альтернатив.\n- Текущий предполагаемый размер исследования не обеспечивает нужную мощность для заложенного эффекта.',
'','## Что нельзя утверждать перед Avito','',
'- Что предотвращены 11,85 реальных уходов, получена фактическая экономия или доказан причинный эффект.\n- Что распределения Monte Carlo оценены по Avito или вероятность провала действительно равна 100%.\n- Что инструмент является медицинской диагностикой или полностью исключает деанонимизацию.\n- Что указанная медиана является реальной ролевой сметой, утверждены владельцы, тарифы или cash бюджет.\n- Что неизвестные юридические, репутационные, knowledge-loss риски и цены альтернатив равны нулю.\n- Что S0, пилот или масштабирование согласованы Avito.',
'','## Открытые содержательные неопределённости','',
'Внутренние baseline и роль-специфические затраты, экономия вакансии и замещающая мощность, реальные ставки и перераспределяемые FTE, цены альтернатив, privacy архитектура, надёжность и валидность сигнала, ICC и размеры команд, причинный outcome, хвосты ущерба и вероятность этапных отказов. Их открытость является частью корректного conditional ТЭО, а не пропущенной подстановкой.',
'','Границы поставленной версии: горизонт фиксирован 60 месяцами; модель до налога, без налогового щита; остаточная стоимость исключена (фиксировано 0), а не оценена как рыночная стоимость IP. Чувствительность и Monte Carlo используют структурную воронку; ITT-режим доступен в сценариях и денежных потоках. Изменение этих границ требует расширения формул и нового QA. Проверка не означает научное, юридическое или внутреннее инвестиционное согласование Avito.',
'','## Manifest и SHA-256','',
'Основные артефакты:',
'',*[f'- `{name}` — `{digest}`' for name,digest in hashes.items()],
'','Неизменённые оригиналы и идентичные рабочие копии:',
'',*[f'- `{p.name}` — `{h}`' for p,h in sources],
'','`manifest.json` содержит версию и SHA-256 всех четырёх итоговых файлов, включая этот аудит. Самохеширование аудита внутри его собственного текста не выполняется.',
]
audit=OUT/'Flicker_TEO_AUDIT.md';audit.write_text('\n'.join(lines)+'\n',encoding='utf-8')
manifest={'version':'1.0','created_at':datetime.now(timezone.utc).isoformat(),'files':[{'name':name,'bytes':(OUT/name).stat().st_size,'sha256':sha(OUT/name)} for name in required+[audit.name]],'source_files':[{'name':p.name,'sha256':h} for p,h in sources],'validation':{'independent_checks':len(results['checks']),'input_checks':len(tests),'formula_errors':0,'sheets':22,'document_pages':len(pdf.pages),'status':'passed'}}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest['validation'],ensure_ascii=False))
