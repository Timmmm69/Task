"""Independent monthly arithmetic, read-only XLSX checks and publication figures."""
import json, math, sys
from statistics import NormalDist
from pathlib import Path
import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter

def spearman(a,b):
    def rank(v):
        _,inverse,counts=np.unique(v,return_inverse=True,return_counts=True)
        end=np.cumsum(counts)
        mean=(end-counts+1+end)/2
        return mean[inverse]
    return float(np.corrcoef(rank(a),rank(b))[0,1])

ROOT=Path(__file__).resolve().parent
QA=ROOT/'qa'; FIG=QA/'figures'; FIG.mkdir(exist_ok=True)
filename=sys.argv[1] if len(sys.argv)>1 else str(QA/'Flicker_TEO_recalculated.xlsx')
w=openpyxl.load_workbook(filename,data_only=True)
f=openpyxl.load_workbook(filename,data_only=False)
mp=json.loads((QA/'workbook_map.json').read_text(encoding='utf-8'))
src=json.loads((QA/'model_data.json').read_text(encoding='utf-8'))
ir=mp['ir']
errors=[]; checks=[]
for s in w:
    for row in s:
        for c in row:
            if c.data_type=='e': errors.append([s.title,c.coordinate,c.value])
checks.append({'test':'Cached formula errors','passed':not errors,'detail':errors[:30]})
maxlen=max(len(c.value) for s in f for row in s for c in row if c.data_type=='f')
checks.append({'test':'Formula length below Excel 8192','passed':maxlen<8192,'detail':maxlen})
checks.append({'test':'No external workbook links','passed':len(f._external_links)==0})

def assert_close(label,a,b,tol=.02):
    ok=bool(isinstance(b,(int,float)) and abs(a-b)<=tol)
    checks.append({'test':label,'passed':ok,'independent':float(a),'workbook':b})
    return ok

def inputs(k): return {key:w['Вводные'].cell(row,4+k).value for key,row in ir.items()}

def forward_gross(net):
    edges=[0,2400000,5000000,20000000,50000000,float('inf')]
    rates=[.13,.15,.18,.20,.22]
    low,high=net,net*2
    for _ in range(90):
        mid=(low+high)/2
        tax=sum(max(0,min(mid,edges[j+1])-edges[j])*r for j,r in enumerate(rates))
        if mid-tax<net: low=mid
        else: high=mid
    return (low+high)/2

def components(p):
    team=[]
    for r in range(7,20):
        gross=forward_gross(w['Команда'].cell(r,3).value*w['Команда']['C23'].value*12)
        taxable=gross*(1+p['bonus'])
        monthly=(taxable+min(taxable,p['contribCap'])*p['contribLow']+max(0,taxable-p['contribCap'])*p['contribHigh']+taxable*p['injury']+gross*p['benefits'])/12
        team.append(monthly)
    loads=np.array([[w['Команда'].cell(r,c).value for c in range(9,15)] for r in range(7,20)])
    team=np.array(team)@loads*p['teamFactor']
    pay=sum(w['Стоимость ухода'].cell(r,3).value*w['Стоимость ухода'].cell(r,4).value for r in range(7,15))
    hourly=sum(w['Стоимость ухода'].cell(r,3).value*w['Стоимость ухода'].cell(r,4).value/w['Стоимость ухода'].cell(r,5).value for r in range(7,15))
    ep={r:w['Стоимость ухода'].cell(r,3).value for r in range(20,28)}
    cash=pay*(ep[20]+ep[24]-ep[22]*p['vacancySavings'])*p['exitFactor']
    other=(p['recruiterHours']*p['recruiterRate']+hourly*(ep[21]+ep[25])+pay*(ep[22]*ep[23]+ep[26]*ep[27])*p['valueMultiple'])*p['exitFactor']
    duration=np.array([2,3,4,12+p['delay'],2,37-p['delay']],dtype=int)
    direct=np.zeros(6)
    for r in range(7,19):
        s=int(w['Прямые расходы'].cell(r,3).value)
        direct[s]+=w['Прямые расходы'].cell(r,6).value/(duration[s] if w['Прямые расходы'].cell(r,4).value=='разово' else 1)
    direct*=1+p['vat']*(1-p['vatRecovery'])
    n=np.array([0,0,p['surveyPilot'],p['surveyExtended'],p['N'],p['N']])
    survey=n*p['participation']*p['surveyFrequency']*p['surveyMinutes']/60/12*hourly*p['timeEnabled']
    manager=p['teams']*n/p['N']*p['signalRate']*p['managerHours']*p['managerRate']*p['timeEnabled']
    training=np.array([0,0,p['teams']*p['surveyPilot']/p['N']*p['trainingHours']*p['managerRate'],0,p['teams']*p['trainingHours']*p['managerRate'],0])*p['timeEnabled']/duration
    time=survey+manager+training
    count=p['N']*p['turnover']*p['regret']*p['address']*p['coverage']*p['participation']*p['privacy']*p['detection']*p['action']*p['causal']
    if p['effectMode']==2: count=p['N']*p['ittARR']
    return dict(team=team,direct=direct,time=time,duration=duration,cash=cash,other=other,count=count,fte=loads.sum(axis=0),survey=survey,manager=manager,training=training)

def model(p,override=None):
    a=components(p)
    if override is not None:
        teamfactor,exitcost,turnover,participation,causal,delay,discount=override
        p=p.copy();p.update(turnover=turnover,participation=participation,causal=causal,delay=int(delay),discount=discount)
        a=components(p);a['team']*=teamfactor
        a['other']=exitcost-a['cash']
    monthly=[]
    stage_ends=np.cumsum(a['duration'])
    for m in range(61):
        if not m: monthly.append([0]*9);continue
        s=int(np.searchsorted(stage_ends,m))
        team=a['team'][s]*(1+p['salaryInflation'])**(m/12)
        time=a['time'][s]*(1+p['salaryInflation'])**(m/12)
        direct=a['direct'][s]*(1+p['infraInflation'])**(m/12)
        teamcash=team*p['teamCashShare'];reserve=(team+direct)*p['reserve']
        cashbenefit=a['count']*a['cash']/12*(1+p['benefitInflation'])**(m/12) if s==5 else 0
        otherbenefit=a['count']*a['other']/12*(1+p['benefitInflation'])**(m/12) if s==5 else 0
        cashnet=cashbenefit-teamcash-direct-(teamcash+direct)*p['reserve']+(p['residual'] if m==60 else 0)
        econnet=cashbenefit+otherbenefit-team-direct-time-reserve+(p['residual'] if m==60 else 0)
        pv=econnet/(1+p['discount'])**(m/12)
        monthly.append([s,cashnet,econnet,pv,team,direct,time,reserve,cashbenefit+otherbenefit])
    cf=np.array(monthly);cost=(a['team']+a['direct'])*(1+p['reserve'])+a['time']
    return dict(p=p,a=a,cf=cf,npv=cf[:,3].sum(),cashnpv=sum(cf[m,1]/(1+p['discount'])**(m/12) for m in range(61)),annualbenefit=a['count']*(a['cash']+a['other']),annualcost=cost[5]*12,phasecost=cost*a['duration'])

models=[model(inputs(k)) for k in range(4)]
for k,m in enumerate(models):
    c=k+3
    for row,key in [(7,'npv'),(8,'cashnpv'),(13,'annualbenefit'),(14,'annualcost')]: assert_close(f'Scenario {k+1} {key}',m[key],w['Сценарии'].cell(row,c).value)
    for j in range(61):
        r=7+k*61+j
        for col,ix in [(14,1),(15,2),(18,3)]: assert_close(f'CF case {k+1} month {j} col {col}',m['cf'][j,ix],w['Денежный поток'].cell(r,col).value)
b=models[1]
assert_close('Sensitivity base equals monthly NPV',b['npv'],w['Чувствительность']['Z23'].value)
z=NormalDist();pw=w['Валидация'];pp=pw['C11'].value
de=1+(pw['C8'].value-1)*pw['C9'].value
ne=pw['C7'].value*(1-pw['C10'].value)/de/2
arr=b['a']['count']/inputs(1)['N'];se=math.sqrt(2*pp*(1-pp)/ne)
critical=z.inv_cdf(1-pw['C12'].value/2);targetz=z.inv_cdf(pw['C13'].value)
assert_close('Power MDE independent NormalDist',(critical+targetz)*se,pw['C19'].value,1e-10)
assert_close('Power sample size independent',math.ceil(4*(critical+targetz)**2*pp*(1-pp)*de/(arr**2*(1-pw['C10'].value))),pw['C21'].value,.01)
assert_close('Power probability independent NormalDist',z.cdf(arr/se-critical)+z.cdf(-arr/se-critical),pw['C22'].value,1e-8)

sens=[]
drivers=[1,b['a']['cash']+b['a']['other'],.08,.7,.3,0,.2]
for i in range(7):
    row=7+i;lo=w['Чувствительность'].cell(row,3).value;hi=w['Чувствительность'].cell(row,5).value
    vv=[]
    for j,v in enumerate([lo,hi]):
        d=drivers.copy();d[i]=v;vnpv=model(inputs(1),d)['npv'];vv.append(vnpv)
        assert_close(f'Sensitivity {i} {j}',vnpv,w['Чувствительность'].cell(row,6+j).value)
    sens.append(dict(name=w['Чувствительность'].cell(row,2).value,low=lo,high=hi,npv_low=vv[0],npv_high=vv[1],swing=abs(vv[1]-vv[0])))
for i in range(5):
    for j in range(5):
        v=drivers.copy();v[1]=w['Пороги'].cell(6,3+j).value;v[4]=w['Пороги'].cell(7+i,2).value
        assert_close(f'Two factor node {i+1} {j+1}',model(inputs(1),v)['npv'],w['Пороги'].cell(7+i,3+j).value)

# Independently generate the fixed-seed uniforms then compute monthly arrays.
# This deliberately does not use the workbook geometric-series formula.
seed=int(w['Вероятностный анализ']['C30'].value);state=seed
uniform=np.empty((10000,5))
for i in range(10000):
    for j in range(5): state=state*16807%2147483647;uniform[i,j]=state/2147483647
draws=np.empty((10000,7))
for j,uindex in enumerate([0,1,2,3,3,0,4]):
    a,mode,c=[w['Вероятностный анализ'].cell(21+j,col).value for col in [3,4,5]]
    u=uniform[:,uindex]
    draws[:,j]=a+np.floor(u*(c-a+1)) if j==5 else np.where(u<=(mode-a)/(c-a),a+np.sqrt(u*(c-a)*(mode-a)),c-np.sqrt((1-u)*(c-a)*(c-mode)))
mc=np.empty((10000,3))
for i,d in enumerate(draws):
    m=model(inputs(1),d);mc[i]=[m['npv'],m['cf'][:37,2].sum(),m['cf'][:,2].sum()]
    if i in [0,1,97,999,4999,9999]:
        for j,col in enumerate([26,27,28]): assert_close(f'MC independent row {i+44} col {col}',mc[i,j],w['Вероятностный анализ'].cell(i+44,col).value,.1)
    if i%2000==0: print(f'Independent simulation {i}/10000',flush=True)
percentiles=np.percentile(mc[:,0],[10,50,90])
excel_mc=np.array([[w['Вероятностный анализ'].cell(44+i,c).value for c in [26,27,28]] for i in range(10000)],dtype=float)
checks.append({'test':'All 30000 MC outcome cells independently agree','passed':bool(np.max(np.abs(excel_mc-mc))<.1),'max_difference':float(np.max(np.abs(excel_mc-mc)))})
excel_draws=np.array([[w['Вероятностный анализ'].cell(44+i,c).value for c in range(2,9)] for i in range(10000)],dtype=float)
checks.append({'test':'All 70000 transformed draws independently agree','passed':bool(np.max(np.abs(excel_draws-draws))<1e-7),'max_difference':float(np.max(np.abs(excel_draws-draws)))})
mcresult=dict(seed=seed,iterations=10000,p10=percentiles[0],p50=percentiles[1],p90=percentiles[2],positive=float(np.mean(mc[:,0]>0)),pay3=float(np.mean(mc[:,1]>=0)),pay5=float(np.mean(mc[:,2]>=0)),expected_loss=float(np.mean(np.maximum(-mc[:,0],0))),downside_at_risk=float(-percentiles[0]),mean=float(mc[:,0].mean()),drivers=[dict(name=sens[j]['name'],spearman=spearman(draws[:,j],mc[:,0])) for j in range(7)])
for row,key in [(8,'p10'),(9,'p50'),(10,'p90'),(11,'positive'),(12,'pay3'),(13,'pay5'),(14,'expected_loss'),(15,'downside_at_risk'),(16,'mean')]: assert_close('MC summary '+key,mcresult[key],w['Вероятностный анализ'].cell(row,3).value,.1)
checks.append({'test':'Monte Carlo 10000 formula rows','passed':sum(f['Вероятностный анализ'].cell(r,26).data_type=='f' for r in range(44,10044))==10000})
np.savez_compressed(QA/'independent_mc.npz',draws=draws,results=mc)
summary=dict(version='1.0',date='2026-09-08',scenario_names=['Консервативный','Базовый','Оптимистичный','Стрессовый'],scenarios=[{key:float(m[key]) for key in ['npv','cashnpv','annualbenefit','annualcost']} for m in models],base=dict(count=b['a']['count'],exitcost=b['a']['cash']+b['a']['other'],exitcash=b['a']['cash'],exitnoncash=b['a']['other'],phasecost=b['phasecost'].tolist(),fte=b['a']['fte'].tolist(),pre_scale=sum(b['phasecost'][:4]),stage0cash=w['Сценарии']['D24'].value,stage0nominal=-b['cf'][1:3,2].sum(),power=w['Валидация']['C22'].value,mde=w['Валидация']['C19'].value,required_n=w['Валидация']['C21'].value,breakeven=w['Сценарии']['D20'].value,rrrbreak=w['Сценарии']['D21'].value,annual_survey=b['a']['survey'][5]*12,annual_manager=b['a']['manager'][5]*12),mc=mcresult,sensitivity=sens,checks=checks,failed=sum(not c['passed'] for c in checks))
(QA/'verified_results.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','figure.facecolor':'white','axes.labelcolor':'#273444','text.color':'#273444'})
navy='#203864';gold='#C39224';red='#A61C32';gray='#8593A5'
def save(name,title,units,draw,height=3.4):
    fig,ax=plt.subplots(figsize=(9,height));draw(ax);ax.set_title(title,loc='left',fontsize=13,pad=15);ax.set_xlabel(units);fig.tight_layout();fig.savefig(FIG/(name+'.png'),dpi=180,bbox_inches='tight');plt.close(fig)
labels=['Нежелательные','Управляемые','Охват','Участие','Privacy','Сигнал','Действие','Предотвращённые']
vals=[w['Удержание'].cell(r,4).value for r in [9,10,11,12,13,14,15,17]]
def funnel(ax):
    ax.barh(labels[::-1],vals[::-1],color=navy);ax.bar_label(ax.containers[0],fmt='%.2f',padding=4);ax.set_xlim(0,370)
save('01_funnel','Удержание возникает только в конце условной воронки','Уходов в год',funnel,4)
st=['Подготовка','Разработка','Пилот','Валидация','Масштаб','Эксплуатация']
def stage(ax):
    a=b['a'];duration=a['duration'];x=np.arange(6);bottom=np.zeros(6)
    for values,label,color in [(a['team']*duration,'Команда',navy),(a['time']*duration,'Время участников',gray),(a['direct']*duration,'Прямые расходы',gold),((a['team']+a['direct'])*inputs(1)['reserve']*duration,'Резерв','#CDD4DD')]:
        ax.bar(x,values/1e6,bottom=bottom,label=label,color=color);bottom+=values/1e6
    ax.set_xticks(x,st);ax.legend(ncol=2,frameon=False);ax.set_ylabel('млн ₽');ax.set_xlabel('Этапы в ценах месяца 0; эксплуатация 37 месяцев')
save('02_stages','Основная стоимость возникает после ранних решений','Этапы',stage)
def effect(ax):
    ax.barh(['Эффект удержания','Полная эксплуатация'],[b['annualbenefit']/1e6,b['annualcost']/1e6],color=[navy,red]);ax.bar_label(ax.containers[0],fmt='%.2f',padding=5);ax.set_xlim(0,118)
save('03_effect','Базовый годовой эффект не покрывает стоимость','млн ₽ в год в ценах месяца 0',effect,2.8)
def cashflow(ax):
    ax.plot(np.cumsum(b['cf'][:,1])/1e6,label='Денежный поток',color=gold);ax.plot(np.cumsum(b['cf'][:,2])/1e6,label='Экономический поток',color=navy);ax.axhline(0,color=gray,lw=.8);ax.legend(frameon=False);ax.set_ylabel('млн ₽');ax.set_xlim(0,60)
save('04_cashflow','За 60 месяцев базовый проект не окупается','Месяц от начала проекта',cashflow)
def scenarios(ax):
    ax.barh(summary['scenario_names'][::-1],[m['npv']/1e6 for m in models][::-1],color=red);ax.set_xlim(-380,0);ax.bar_label(ax.containers[0],fmt='%.1f',padding=3);ax.axvline(0,color=gray)
save('05_scenarios','Во всех четырёх сценариях экономический NPV отрицателен','NPV за 60 месяцев млн ₽',scenarios)
ordered=sorted(sens,key=lambda x:x['swing'])
def tornado(ax):
    for i,s in enumerate(ordered):
        lo,hi=sorted([s['npv_low']/1e6,s['npv_high']/1e6]);ax.barh(i,hi-lo,left=lo,color=navy,height=.6)
    ax.set_yticks(range(7),[s['name'] for s in ordered]);ax.axvline(b['npv']/1e6,color=gold,ls='--',label='Базовый NPV');ax.legend(frameon=False)
save('06_tornado','Наиболее важны диапазоны с наибольшим размахом NPV','Экономический NPV млн ₽; по одному драйверу',tornado,4)
def breakeven(ax):
    x=np.linspace(0,112,101);ax.plot(x,x*summary['base']['exitcost']/1e6,color=navy,label='Экономический эффект');ax.axhline(b['annualcost']/1e6,color=red,label='Годовые затраты');ax.set_xlim(0,112);ax.set_ylabel('млн ₽ в год');ax.legend(frameon=False)
save('07_breakeven','Даже весь управляемый пул не покрывает базовую стоимость','Предотвращённых уходов в год; доступный пул 112',breakeven)
riskrows=['Валидность сигнала','Мощность','Причинный эффект','Участие и смещение','Приватность и вред','Готовность HR','Интеграции и объём','Правовые требования']
presence=np.array([[0,1,1,1,1,1],[1,1,1,1,0,0],[0,0,1,1,1,1],[0,0,1,1,1,1],[1,1,1,1,1,1],[1,0,1,1,1,1],[1,1,0,1,1,1],[1,1,1,1,1,1]])
def risk(ax):
    for i in range(8):
        for j in range(6):
            if presence[i,j]: ax.scatter(j,i,color=navy,s=65)
    ax.set_yticks(range(8),riskrows);ax.set_xticks(range(6),st);ax.invert_yaxis();ax.set_xlim(-.5,5.5);ax.grid(axis='x',alpha=.2)
save('08_risks','Карта этапов возникновения рисков без вымышленных вероятностей','Точка означает необходимость контроля; тяжесть и вероятность неизвестны',risk,4)
def timeline(ax):
    durations=b['a']['duration'];begins=np.r_[0,np.cumsum(durations)[:-1]]
    ax.barh(st[::-1],durations[::-1],left=begins[::-1],color=navy);ax.set_xlim(0,60)
    for x in np.cumsum(durations)[:-1]: ax.axvline(x,color=gold,ls=':',lw=1)
save('09_timeline','Каждый этап требует отдельного решения о продолжении','Месяцы от старта; линии означают stop go',timeline,3.5)
def distribution(ax):
    ax.hist(mc[:,0]/1e6,bins=45,color=navy,edgecolor='white',linewidth=.3);ax.axvline(percentiles[1]/1e6,color=gold,label='P50');ax.set_ylabel('Итераций');ax.legend(frameon=False)
save('10_montecarlo','Иллюстративное распределение не является прогнозом Avito','Экономический NPV млн ₽',distribution)
print(json.dumps({'failed':summary['failed'],'mc':mcresult,'base':summary['base']},ensure_ascii=False,indent=2))
